from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Bao_cao_WordPress_Quang_ba_San_pham.docx"
BLUE = "1F4E79"
LIGHT_BLUE = "D9EAF7"
GRAY = "F2F4F5"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(10)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Trang ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def add_placeholder(document, caption):
    table = document.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Cm(15.5)
    cell = table.cell(0, 0)
    set_cell_shading(cell, GRAY)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(20)
    paragraph.paragraph_format.space_after = Pt(20)
    run = paragraph.add_run(caption)
    run.italic = True
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(90, 100, 110)


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.add_run(text)


def add_body(document, text):
    paragraph = document.add_paragraph(text)
    paragraph.paragraph_format.first_line_indent = Cm(0.7)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.12
    return paragraph


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for index, text in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], text, bold=True, color=BLUE)
        set_cell_shading(table.rows[0].cells[index], LIGHT_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for index, text in enumerate(row):
            set_cell_text(cells[index], text)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def make_document():
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(1)
    section.footer_distance = Cm(1)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    for style_name, size in (("Title", 25), ("Heading 1", 17), ("Heading 2", 13)):
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(BLUE)
    styles["Heading 1"].font.bold = True
    styles["Heading 2"].font.bold = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("TRIỂN KHAI VÀ QUẢN TRỊ HỆ THỐNG PHẦN MỀM")
    header_run.font.name = "Times New Roman"
    header_run.font.size = Pt(8)
    header_run.font.color.rgb = RGBColor(100, 110, 120)
    add_page_number(section.footer.paragraphs[0])

    # Cover page
    for _ in range(2):
        document.add_paragraph()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TRƯỜNG: ........................................................")
    r.bold = True
    r.font.size = Pt(14)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("KHOA / BỘ MÔN: ................................................").font.size = Pt(13)
    for _ in range(2):
        document.add_paragraph()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BÁO CÁO ĐỀ TÀI THỰC HÀNH")
    r.bold = True
    r.font.size = Pt(19)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TRIỂN KHAI VÀ QUẢN TRỊ HỆ THỐNG PHẦN MỀM")
    r.bold = True
    r.font.size = Pt(17)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    r = p.add_run("WEBSITE QUẢNG BÁ SẢN PHẨM\nTRÊN NỀN TẢNG WORDPRESS")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    for _ in range(2):
        document.add_paragraph()
    add_table(document, ["Thông tin", "Nội dung"], [
        ("Sinh viên", "[ĐIỀN HỌ VÀ TÊN]"),
        ("Mã sinh viên", "[ĐIỀN MSSV]"),
        ("Lớp", "[ĐIỀN LỚP]"),
        ("Giảng viên", "[ĐIỀN TÊN GIẢNG VIÊN]"),
    ])
    p = document.add_paragraph("[ĐỊA DANH], tháng 10 năm 2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()

    pages = [
        ("1. TÓM TẮT VÀ MỤC TIÊU", [
            ("h", "Tóm tắt đề tài"),
            ("p", "Đề tài xây dựng một website quảng bá sản phẩm động bằng WordPress official image, lưu nội dung trong MySQL và vận hành toàn bộ bằng Docker Compose. Nginx làm reverse proxy, phpMyAdmin hỗ trợ quản trị cơ sở dữ liệu, Prometheus thu thập metrics, Grafana hiển thị dashboard, còn Loki và Promtail tập trung log container. Website không phải trang HTML tĩnh: quản trị viên có thể tạo, sửa và phân loại sản phẩm qua WordPress Admin; nội dung được WordPress truy vấn từ database."),
            ("h", "Mục tiêu"),
            ("b", "Triển khai ứng dụng có danh mục sản phẩm, tìm kiếm, trang chi tiết, giá, ảnh và liên kết liên hệ/mua hàng."),
            ("b", "Đóng gói các dịch vụ thành cấu hình tái lập được bằng Docker Compose và lưu dữ liệu bền vững bằng named volumes."),
            ("b", "Tổ chức mạng, giới hạn cổng truy cập, quản lý bí mật qua biến môi trường và áp dụng security headers."),
            ("b", "Quan sát trạng thái container, web server, database và log tập trung để hỗ trợ vận hành, phát hiện lỗi."),
            ("h", "Phạm vi báo cáo"),
            ("p", "Báo cáo mô tả kiến trúc, cấu hình, quy trình triển khai trên Ubuntu VM và các kiểm tra đã thực hiện trong quá trình chạy thử. Những trường thông tin cá nhân và ảnh chụp minh chứng được đánh dấu để sinh viên bổ sung trước khi nộp."),
        ]),
        ("2. YÊU CẦU VÀ TIÊU CHÍ HOÀN THÀNH", [
            ("p", "Hệ thống được xây dựng theo yêu cầu đề tài Website Quảng bá Sản phẩm (WordPress). Các yêu cầu chức năng và vận hành được ánh xạ trực tiếp với dịch vụ trong Compose như sau:"),
            ("table", (["Yêu cầu", "Thành phần đáp ứng", "Kiểm chứng"], [
                ("Website động", "WordPress, theme Product Showcase, MySQL", "Trang /san-pham/ hiển thị sản phẩm và danh mục"),
                ("DB và giao diện quản trị", "MySQL 8.4, phpMyAdmin", "Đăng nhập phpMyAdmin qua localhost:8081"),
                ("Reverse proxy", "Nginx 1.28 Alpine", "Truy cập website qua cổng 80"),
                ("Metrics", "Prometheus, Grafana, cAdvisor, exporters", "Các scrape targets báo UP"),
                ("Log tập trung", "Loki và Promtail", "Log hiển thị trong Grafana Explore/dashboard"),
                ("Hardening", "Networks riêng, secrets trong .env, headers, non-root", "Compose và cấu hình Nginx"),
            ])),
            ("h", "Tiêu chí nghiệm thu"),
            ("b", "Các service khởi động ổn định; WordPress đọc/ghi database và có dữ liệu sản phẩm."),
            ("b", "Nginx, MySQL và cAdvisor được scrape; logs có thể truy vấn bằng LogQL."),
            ("b", "MySQL không publish cổng ra host; cổng quản trị bind vào loopback."),
            ("p", "Tại thời điểm tổng hợp báo cáo, ảnh Prometheus Targets cho thấy cAdvisor, Loki, MySQL, Nginx, Prometheus và Promtail đều UP sau khi sửa cấu hình MySQL exporter 0.17. Dashboard đã provision; cần chèn thêm ảnh dashboard sau khi refresh để xác nhận các biểu đồ metrics hiển thị dữ liệu."),
        ]),
        ("3. KIẾN TRÚC TỔNG THỂ", [
            ("p", "Hệ thống gồm các mặt phẳng ứng dụng, dữ liệu, giám sát và logging. Trình duyệt gửi HTTP đến Nginx ở cổng 80. Nginx chuyển tiếp request web đến WordPress; WordPress kết nối MySQL qua mạng backend. phpMyAdmin cũng kết nối MySQL nhưng chỉ công bố trên loopback host. Prometheus truy vấn exporters; Grafana sử dụng Prometheus và Loki làm datasource."),
            ("table", (["Mạng", "Thành viên chính", "Mục đích"], [
                ("frontend", "Nginx, WordPress", "Đường đi request web"),
                ("backend (internal)", "WordPress, MySQL, phpMyAdmin, MySQL exporter/config", "Database không có đường ra host"),
                ("monitoring", "Prometheus, Grafana, cAdvisor, exporters, Loki, Promtail", "Thu thập và truy vấn metrics/logs"),
                ("logging (internal)", "Loki, Promtail", "Đường trao đổi log nội bộ"),
            ])),
            ("h", "Luồng request và dữ liệu"),
            ("b", "Browser → cổng host 80 → Nginx:8080 → WordPress:80 → MySQL:3306."),
            ("b", "phpMyAdmin: host 127.0.0.1:8081 → phpMyAdmin:80 → MySQL:3306."),
            ("b", "Prometheus scrape cAdvisor:8080, nginx-exporter:9113, mysqld-exporter:9104, Loki:3100 và Promtail:9080."),
            ("b", "Promtail đọc Docker JSON logs từ host và push đến Loki; Grafana truy vấn Loki qua datasource nội bộ."),
            ("p", "Docker Compose cung cấp DNS theo tên service trong từng network. Một service chỉ giao tiếp được khi có network chung; vì vậy network membership được thiết kế theo luồng cần thiết thay vì nối mọi container vào một mạng phẳng."),
        ]),
        ("4. WEBSITE WORDPRESS VÀ CHỨC NĂNG SẢN PHẨM", [
            ("p", "Ứng dụng dùng image chính thức wordpress:php8.3-apache. Mã lõi WordPress, plugin và dữ liệu runtime nằm trong named volume wordpress_data; theme Product Showcase được mount read-only từ thư mục dự án. Database dùng MySQL 8.4 với volume mysql_data. Cách tổ chức này giữ dữ liệu qua lần recreate container và cho phép quản lý theme trong Git."),
            ("h", "Mô hình nội dung"),
            ("b", "Custom post type product lưu tên, mô tả ngắn, nội dung chi tiết và trạng thái xuất bản."),
            ("b", "Taxonomy product_category hỗ trợ phân loại sản phẩm theo nhóm."),
            ("b", "Metadata sản phẩm lưu giá, URL liên hệ/mua hàng và URL ảnh dự phòng; ảnh đại diện WordPress được ưu tiên nếu đã tải lên."),
            ("b", "Các giao diện gồm trang chủ, danh sách /san-pham/, lọc danh mục, tìm kiếm và trang chi tiết."),
            ("h", "Quản trị và dữ liệu demo"),
            ("p", "Theme có meta box để biên tập giá, URL và ảnh. Ba sản phẩm demo được tạo tự động khi theme đã bật và request đầu tiên chạy hook khởi tạo. Đây là dữ liệu mẫu; sinh viên cần thay mô tả, ảnh và thông tin bán hàng bằng nội dung phù hợp đề tài hoặc ghi rõ đây là dữ liệu minh họa."),
            ("p", "Trong phiên chạy thử, trang /san-pham/ đã hiển thị ba sản phẩm, ba danh mục và ô tìm kiếm. Có thể bổ sung menu, trang Giới thiệu, thông tin liên hệ và chính sách để hoàn thiện trải nghiệm website quảng bá."),
            ("placeholder", "[CHÈN ẢNH 1: TRANG DANH SÁCH SẢN PHẨM]"),
        ]),
        ("5. DOCKER COMPOSE, DATABASE VÀ VOLUME", [
            ("p", "Tệp docker-compose.yml khai báo toàn bộ dịch vụ, restart policy, healthcheck, network, volumes, port mapping và secret interpolation. Điều này giúp triển khai nhất quán thay vì cấu hình thủ công từng container. Compose yêu cầu các biến bắt buộc trong .env; file mẫu .env.example không chứa secret thật và .env được loại khỏi Git bằng .gitignore."),
            ("h", "Database và quyền truy cập"),
            ("b", "MYSQL_DATABASE và MYSQL_USER tạo database cùng user ứng dụng; WordPress sử dụng user ứng dụng, không đăng nhập root."),
            ("b", "MYSQL_ROOT_PASSWORD chỉ dùng cho quản trị khởi tạo/khẩn cấp; cổng MySQL không publish ra host."),
            ("b", "Healthcheck mysqladmin ping giúp WordPress/phpMyAdmin chờ database sẵn sàng."),
            ("b", "phpMyAdmin publish 127.0.0.1:8081:80 và chỉ tham gia backend; không đi qua Nginx public."),
            ("h", "Tính bền vững và exporter"),
            ("p", "Named volumes lưu database, WordPress files, Prometheus TSDB, Grafana data, Loki data và Promtail positions. Dịch vụ mysql-exporter-config dùng Alpine, tạo .my.cnf từ biến môi trường trong volume riêng rồi thoát thành công. mysqld-exporter 0.17 mount file read-only; đây là cách phù hợp với cấu hình phiên bản đã triển khai."),
            ("p", "Lệnh docker compose down chỉ dừng và xóa container/network; docker compose down -v xóa cả named volumes và làm mất dữ liệu. Không dùng -v khi không có bản sao lưu."),
        ]),
        ("6. NGINX REVERSE PROXY VÀ HTTP SECURITY", [
            ("p", "Nginx 1.28 Alpine là điểm vào duy nhất của website ở cổng host 80, chuyển tiếp request đến WordPress qua frontend network. Cấu hình nginx.conf đặt PID và các thư mục temporary dưới /tmp để hỗ trợ user không phải root cùng filesystem read-only. Nginx chạy UID 101, bỏ toàn bộ Linux capabilities và dùng tmpfs giới hạn cho /tmp."),
            ("h", "Headers và giới hạn truy cập"),
            ("b", "X-Frame-Options: SAMEORIGIN hạn chế nhúng trang từ origin khác."),
            ("b", "X-Content-Type-Options: nosniff giảm diễn giải sai MIME type."),
            ("b", "Referrer-Policy và Permissions-Policy hạn chế thông tin referrer và một số quyền trình duyệt."),
            ("b", "server_tokens off giảm tiết lộ phiên bản Nginx; client_max_body_size giới hạn upload."),
            ("b", "Nginx stub_status nằm trên listener nội bộ port 8081, không nằm trên listener website 8080; exporter scrape qua monitoring network."),
            ("h", "Giới hạn triển khai"),
            ("p", "Stack hiện dùng HTTP trong môi trường lab, nhưng đáp ứng lựa chọn security headers cơ bản trong yêu cầu. Security headers không mã hóa dữ liệu như TLS. Nếu công bố website ra Internet cần HTTPS tại Nginx hoặc một TLS reverse proxy tin cậy, cấu hình domain/certificate, HSTS chỉ sau khi HTTPS ổn định, và rà soát CSP phù hợp với WordPress/plugin."),
            ("placeholder", "[CHÈN ẢNH 2: WEBSITE WORDPRESS QUA NGINX VÀ SECURITY HEADERS]"),
        ]),
        ("7. PROMETHEUS, EXPORTERS VÀ GRAFANA", [
            ("p", "Prometheus 3.5.0 scrape theo chu kỳ 15 giây, lưu TSDB trong volume riêng và đặt retention 15 ngày. Cấu hình có các job cho chính Prometheus, cAdvisor, Nginx exporter, MySQL exporter, Loki và Promtail. cAdvisor cung cấp metrics container; Nginx exporter lấy stub_status; MySQL exporter truy vấn trạng thái và biến hệ thống của database."),
            ("h", "Dashboard"),
            ("p", "Grafana 12.1.1 dùng tài khoản quản trị từ .env, tắt đăng ký người dùng, provision datasource Prometheus/Loki bằng UID cố định và nạp dashboard WordPress Product Site từ JSON. Dashboard có panel CPU container, working-set memory, kết nối và request rate của Nginx, kết nối MySQL, trạng thái exporter và log gần đây."),
            ("h", "Xác minh scrape"),
            ("p", "Ban đầu mysqld-exporter báo lỗi do v0.17 không đọc cấu hình theo DATA_SOURCE_NAME như biến Compose cũ. Sau khi chuyển sang .my.cnf trong volume riêng và recreate service, trang Prometheus Targets được chụp lại với các job cAdvisor, Loki, MySQL, Nginx, Prometheus và Promtail ở trạng thái UP. Đây là minh chứng cho luồng scrape; biểu đồ cần được refresh sau khi có mẫu dữ liệu."),
            ("placeholder", "[CHÈN ẢNH 3: PROMETHEUS TARGETS TẤT CẢ UP]"),
            ("placeholder", "[CHÈN ẢNH 4: GRAFANA DASHBOARD SAU KHI REFRESH, CÓ METRICS]"),
        ]),
        ("8. LOKI, PROMTAIL VÀ TRUY VẤN LOGQL", [
            ("p", "Loki 3.5.3 lưu log dạng filesystem với retention 7 ngày. Promtail 3.5.3 phát hiện container qua Docker service discovery, chỉ giữ container thuộc Compose project wordpress-product-site, đọc Docker JSON log và gắn labels container/service trước khi gửi đến Loki. Vị trí đọc log được lưu trong named volume để tránh đọc lại toàn bộ dữ liệu sau khi container khởi động lại."),
            ("h", "LogQL minh họa"),
            ("code", '{service="wordpress"}'),
            ("code", '{service="nginx"} |= " 404 "'),
            ("code", '{container=~".*(mysql|phpmyadmin).*"} |~ "(?i)error|warning"'),
            ("p", "Truy vấn đầu lấy log WordPress; truy vấn thứ hai lọc request 404 của Nginx; truy vấn thứ ba tìm lỗi/cảnh báo trong log database và phpMyAdmin. Label cần được giữ có cardinality hợp lý; không gắn ID request hoặc nội dung thay đổi liên tục thành label."),
            ("h", "Bảo vệ log"),
            ("p", "Promtail cần đọc Docker socket và log host nên chạy với quyền cao; Docker socket tương đương năng lực quản trị Docker. Chỉ dùng cấu hình này trên host lab tin cậy, không mount socket vào container ứng dụng. Logging network là internal; giao diện Grafana chỉ bind vào loopback."),
            ("placeholder", "[CHÈN ẢNH 5: LOG WORDPRESS/NGINX TRONG GRAFANA EXPLORE]"),
        ]),
        ("9. HARDENING VÀ RỦI RO CÒN LẠI", [
            ("h", "Biện pháp đã áp dụng"),
            ("b", "Tách frontend, backend, monitoring và logging; backend/logging dùng internal network."),
            ("b", "Dùng biến môi trường cho mật khẩu; .env không được commit; dịch vụ cần secret sử dụng biến bắt buộc khi Compose render."),
            ("b", "MySQL không publish cổng; phpMyAdmin, Grafana, Prometheus chỉ bind trên loopback host."),
            ("b", "WordPress kết nối bằng user ứng dụng; MySQL exporter dùng user riêng; phpMyAdmin không truy cập arbitrary server."),
            ("b", "Nginx, Prometheus, Grafana, Loki và MySQL exporter chạy non-root; Nginx dùng read-only filesystem và drop capabilities."),
            ("b", "WordPress tắt chỉnh sửa theme/plugin trực tiếp từ Dashboard qua DISALLOW_FILE_EDIT; Nginx thêm các HTTP security headers."),
            ("h", "Trade-off và rủi ro"),
            ("p", "MySQL exporter cần quyền SELECT rộng cùng PROCESS/REPLICATION CLIENT cho các collector mặc định; tài khoản chỉ đọc nhưng vẫn có thể đọc dữ liệu ứng dụng. cAdvisor cần privileged và mount host. Promtail cần quyền root/socket. Các quyền này là yêu cầu kỹ thuật của việc thu thập metrics/log trong cấu hình hiện tại, không nên xem là hardening tuyệt đối."),
            ("p", "Trước khi triển khai Internet cần TLS, firewall theo nguyên tắc tối thiểu, secret manager, backup/restore đã thử nghiệm, giới hạn quyền exporter nếu thay collector, rà image/update policy, rate limiting và kiểm thử bảo mật phù hợp. Không chạy phpMyAdmin công khai."),
        ]),
        ("10. CÀI ĐẶT, VẬN HÀNH VÀ SAO LƯU", [
            ("h", "Môi trường triển khai đã kiểm tra"),
            ("p", "Stack được chạy trong Ubuntu 26.04 LTS trên VMware Workstation, Docker Engine 29.9.2 và Docker Compose 5.6.1. Docker info ghi nhận 2 CPU và khoảng 2.5 GiB RAM. Toàn bộ images được pull và các container chạy; với số tài nguyên này cần theo dõi bộ nhớ nếu dashboard chậm hoặc service bị OOM."),
            ("h", "Quy trình chạy"),
            ("code", "cp .env.example .env\nchmod 600 .env\ndocker compose config\ndocker compose up -d\ndocker compose ps"),
            ("h", "Quy trình quản trị thường ngày"),
            ("code", "docker compose logs -f wordpress nginx\ndocker compose restart wordpress\ndocker compose ps"),
            ("p", "WordPress được cài đặt qua trình duyệt ở http://localhost/, sau đó bật Product Showcase trong Appearance > Themes. Sản phẩm quản lý ở menu Products; trang danh mục là /san-pham/. phpMyAdmin ở localhost:8081; Grafana ở localhost:3000; Prometheus ở localhost:9090. Các cổng quản trị chỉ truy cập từ loopback của Ubuntu VM."),
            ("h", "Sao lưu"),
            ("p", "Sao lưu cần bao gồm database MySQL và nội dung WordPress uploads; có thể bổ sung volume Grafana/Prometheus/Loki tùy yêu cầu phục hồi. Chỉ sao lưu volume không thay thế quy trình kiểm thử restore. docker compose down -v sẽ xóa dữ liệu và không dùng trong vận hành thường ngày."),
        ]),
        ("11. KIỂM THỬ VÀ KẾT QUẢ", [
            ("table", (["Mã", "Kiểm thử", "Kết quả ghi nhận"], [
                ("T01", "Compose render cấu hình", "Thành công: docker compose config --quiet không báo lỗi."),
                ("T02", "Khởi động dịch vụ", "Thành công: images/volumes/networks được tạo, WordPress và MySQL Healthy."),
                ("T03", "Website sản phẩm", "Thành công: /san-pham/ hiển thị ba sản phẩm mẫu, danh mục và tìm kiếm."),
                ("T04", "Prometheus Targets", "Thành công sau sửa exporter: cAdvisor, Loki, MySQL, Nginx, Prometheus, Promtail đều UP."),
                ("T05", "Loki/Promtail", "Thành công: log ứng dụng hiển thị trong panel Grafana."),
                ("T06", "Grafana metrics panels", "Dashboard đã provision; cần lưu ảnh sau refresh để chứng minh giá trị panel."),
                ("T07", "phpMyAdmin", "Được cấu hình bind loopback :8081; bổ sung ảnh đăng nhập/truy cập nếu tiêu chí chấm yêu cầu."),
            ])),
            ("h", "Nhận xét"),
            ("p", "Kết quả xác nhận đường đi web, database healthcheck, scrape targets và log ingestion. Phần truy vấn MySQL exporter được khắc phục bằng file cấu hình .my.cnf; ảnh Targets cuối cùng cho thấy job MySQL UP. Cần chụp lại dashboard sau sửa và refresh để chứng minh time series có dữ liệu, không dùng ảnh dashboard No data trước đó làm minh chứng hoàn tất."),
            ("placeholder", "[CHÈN ẢNH 6: docker compose ps VÀ TRẠNG THÁI CONTAINER]"),
            ("placeholder", "[CHÈN ẢNH 7: GRAFANA DASHBOARD SAU KHI TARGETS UP]"),
        ]),
        ("12. QUẢN LÝ MÃ NGUỒN, KẾT LUẬN VÀ TÀI LIỆU", [
            ("h", "Kế hoạch ba commit"),
            ("b", "Commit 1: WordPress, MySQL, phpMyAdmin, theme ứng dụng và Nginx reverse proxy/security headers."),
            ("b", "Commit 2: Prometheus, Grafana, cAdvisor/exporters và provisioning dashboard/datasource."),
            ("b", "Commit 3: Loki, Promtail, logging và hardening; cập nhật README, MySQL exporter config."),
            ("p", "Các commit phải được tạo tuần tự trên GitHub để lịch sử phản ánh từng giai đoạn; nếu chỉ commit trạng thái cuối, lịch sử sẽ không có ba mốc như yêu cầu. Trước khi push, kiểm tra git status và bảo đảm .env không nằm trong staged files."),
            ("h", "Kết luận"),
            ("p", "Đề tài triển khai thành công một website WordPress động có database, reverse proxy, metrics và logging tập trung trên Docker Compose. Hệ thống phân tách mạng và giới hạn một số cổng quản trị, đồng thời cung cấp quy trình quan sát bằng Grafana và Prometheus. Quá trình kiểm thử cũng cho thấy exporter cần cấu hình khớp đúng phiên bản; việc đọc target health và log giúp xác định, sửa lỗi vận hành."),
            ("p", "Các hướng phát triển gồm bổ sung HTTPS/domain, backup tự động, alert rules, dashboard theo nhu cầu nghiệp vụ, nội dung sản phẩm thực tế và giới hạn quyền exporter. Trước khi nộp, sinh viên điền thông tin cá nhân, chèn ảnh minh chứng mới nhất và rà soát ba commit."),
            ("h", "Tài liệu tham khảo"),
            ("b", "Docker Compose documentation: https://docs.docker.com/compose/"),
            ("b", "WordPress Developer Resources: https://developer.wordpress.org/"),
            ("b", "Prometheus documentation: https://prometheus.io/docs/"),
            ("b", "Grafana, Loki, Promtail documentation: https://grafana.com/docs/"),
            ("b", "MySQL 8.4 Reference Manual: https://dev.mysql.com/doc/refman/8.4/en/"),
        ]),
    ]

    for page_index, (title, blocks) in enumerate(pages):
        document.add_heading(title, level=1)
        for kind, payload in blocks:
            if kind == "h":
                document.add_heading(payload, level=2)
            elif kind == "p":
                add_body(document, payload)
            elif kind == "b":
                add_bullet(document, payload)
            elif kind == "code":
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.5)
                paragraph.paragraph_format.space_after = Pt(7)
                run = paragraph.add_run(payload)
                run.font.name = "Consolas"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
                run.font.size = Pt(9)
            elif kind == "table":
                headers, rows = payload
                add_table(document, headers, rows)
            elif kind == "placeholder":
                add_placeholder(document, payload)
        if page_index < len(pages) - 1:
            document.add_page_break()

    document.core_properties.title = "Báo cáo Website quảng bá sản phẩm bằng WordPress"
    document.core_properties.subject = "Triển khai và Quản trị Hệ thống Phần mềm"
    document.core_properties.author = "[ĐIỀN HỌ VÀ TÊN SINH VIÊN]"
    document.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    make_document()